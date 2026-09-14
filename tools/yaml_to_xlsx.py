#!/usr/bin/env python3
"""Flatten docs/phase2/rules.yaml into bots/*/Data/rules.xlsx.

UiPath reads configuration with Read Range, not a YAML parser. So the YAML is
the editable source of truth (reviewable by the coordinator who owns the
policy) and this script produces the runtime form the bots actually load.

    python tools/yaml_to_xlsx.py                  # write to all three bots
    python tools/yaml_to_xlsx.py --out path.xlsx  # write one file

Sheet "Rules": Key | Value | Type | Section
Keys are dotted paths, e.g. `allocation.weights.weight_earliest`, so a bot
reads one flat lookup dictionary and indexes it by the same names used in
docs/phase2/bot_scheduling.md.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
RULES_YAML = REPO_ROOT / "docs" / "phase2" / "rules.yaml"
BOT_DIRS = ["SchedulingBot", "ReminderBot", "ConflictResolutionBot"]


def flatten(node, prefix="") -> list[tuple[str, object]]:
    """Dotted-path flatten. Lists of scalars become a pipe-joined string;
    lists of dicts are indexed, so `passes[0].name` stays addressable."""
    out = []
    if isinstance(node, dict):
        for k, v in node.items():
            out.extend(flatten(v, f"{prefix}.{k}" if prefix else str(k)))
    elif isinstance(node, list):
        if node and all(not isinstance(x, (dict, list)) for x in node):
            out.append((prefix, "|".join(str(x) for x in node)))
        else:
            for i, item in enumerate(node):
                out.extend(flatten(item, f"{prefix}[{i}]"))
    else:
        out.append((prefix, node))
    return out


def type_name(v) -> str:
    if isinstance(v, bool):
        return "Boolean"
    if isinstance(v, int):
        return "Int32"
    if isinstance(v, float):
        return "Double"
    return "String"


def to_cell(v):
    """UiPath's Read Range gives strings back; booleans must be unambiguous."""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if v is None:
        return ""
    return v


def build_workbook(rules: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"

    headers = ["Key", "Value", "Type", "Section"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F5F8B")
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")

    rows = flatten(rules)
    for key, value in rows:
        ws.append([key, to_cell(value), type_name(value), key.split(".")[0]])

    widths = [58, 42, 10, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # A second sheet the bots do not read — it exists so a human opening the
    # workbook can see where it came from and not edit the wrong artefact.
    meta = wb.create_sheet("README")
    meta["A1"] = "Generated file — do not edit by hand"
    meta["A1"].font = Font(bold=True, size=12)
    for r, line in enumerate(
        [
            "",
            "Source:    docs/phase2/rules.yaml",
            "Generated: python tools/yaml_to_xlsx.py",
            "",
            "Edit the YAML and regenerate. Changes made directly here are lost",
            "on the next run and will silently diverge from what the HMS uses,",
            "which reads the YAML directly.",
            "",
            f"Rules version: {rules.get('meta', {}).get('version', 'unknown')}",
            f"Effective from: {rules.get('meta', {}).get('effective_from', 'unknown')}",
            f"Keys exported: {len(rows)}",
        ],
        start=2,
    ):
        meta[f"A{r}"] = line
    meta.column_dimensions["A"].width = 70
    return wb


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Generate rules.xlsx from rules.yaml")
    ap.add_argument("--src", default=str(RULES_YAML))
    ap.add_argument("--out", help="single output path (default: all three bot Data folders)")
    args = ap.parse_args(argv)

    src = Path(args.src)
    if not src.exists():
        print(f"error: {src} not found", file=sys.stderr)
        return 1

    with open(src, "r", encoding="utf-8") as fh:
        rules = yaml.safe_load(fh)

    targets = (
        [Path(args.out)]
        if args.out
        else [REPO_ROOT / "bots" / b / "Data" / "rules.xlsx" for b in BOT_DIRS]
    )

    n = len(flatten(rules))
    for t in targets:
        t.parent.mkdir(parents=True, exist_ok=True)
        build_workbook(rules).save(t)
        print(f"wrote {n} rules -> {t.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

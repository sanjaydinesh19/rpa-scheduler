#!/usr/bin/env python3
"""Create and reconcile data/appointments.xlsx — the second "system".

The Scheduling Bot appends to this workbook as part of each booking
transaction; the Conflict Bot reconciles it against the HMS on a sweep. This
script does three jobs:

    python tools/excel_mirror.py init        # create the empty workbook
    python tools/excel_mirror.py sync        # pull unsynced HMS rows into it
    python tools/excel_mirror.py reconcile   # report HMS vs Excel differences

Direction of truth is one-way on purpose: the HMS is the system of record and
Excel gets corrected. A row in Excel with no HMS appointment is reported, never
deleted — it may be the only trace of a booking the HMS lost.
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = REPO_ROOT / "data" / "appointments.xlsx"
SHEET = "Appointments"

COLUMNS = [
    "ReferenceNo", "PatientMRN", "PatientName", "Phone", "Department",
    "DoctorName", "AppointmentDate", "AppointmentTime", "Status",
    "Priority", "BookedBy", "SyncedAt",
]


def init(path: Path = EXCEL_PATH) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    fill = PatternFill("solid", fgColor="1F5F8B")
    for i in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    widths = [18, 12, 22, 16, 18, 24, 14, 14, 12, 10, 16, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_rows(path: Path = EXCEL_PATH) -> dict[str, dict]:
    if not path.exists():
        return {}
    ws = load_workbook(path)[SHEET]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        out[str(row[0])] = dict(zip(COLUMNS, row))
    return out


def append_row(appt: dict, path: Path = EXCEL_PATH) -> None:
    if not path.exists():
        init(path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    dt = datetime.fromisoformat(appt["appointment_datetime"])
    p = appt.get("patient") or {}
    ws.append([
        appt["reference_no"], p.get("mrn", ""), p.get("full_name", ""), p.get("phone", ""),
        appt.get("department", ""), appt.get("doctor_name", ""),
        dt.date().isoformat(), dt.strftime("%H:%M"),
        appt["status"], appt["priority"], appt["booked_by"],
        datetime.now().isoformat(timespec="seconds"),
    ])
    wb.save(path)


def api(base: str, key: str, path_: str, **params):
    r = requests.get(f"{base.rstrip('/')}/api/v1{path_}",
                     headers={"X-API-Key": key}, params=params, timeout=30)
    r.raise_for_status()
    return r.json()


def sync(base: str, key: str, path: Path = EXCEL_PATH) -> dict:
    appts = api(base, key, "/appointments", include_patient="true", limit=1000)
    existing = read_rows(path)
    added = 0
    for a in appts:
        if a["reference_no"] in existing:
            continue
        append_row(a, path)
        requests.patch(
            f"{base.rstrip('/')}/api/v1/appointments/{a['appointment_id']}",
            headers={"X-API-Key": key},
            json={"synced_to_excel_at": True},
            timeout=30,
        )
        added += 1
    return {"hms_rows": len(appts), "already_present": len(existing), "added": added}


def reconcile(base: str, key: str, path: Path = EXCEL_PATH) -> dict:
    appts = {a["reference_no"]: a for a in api(base, key, "/appointments", limit=1000)}
    excel = read_rows(path)

    findings = {"MISSING_IN_EXCEL": [], "STATUS_DRIFT": [],
                "DATETIME_DRIFT": [], "ORPHAN_IN_EXCEL": []}

    for ref, a in appts.items():
        e = excel.get(ref)
        if e is None:
            findings["MISSING_IN_EXCEL"].append(ref)
            continue
        if str(e["Status"]) != a["status"]:
            findings["STATUS_DRIFT"].append(
                {"ref": ref, "excel": e["Status"], "hms": a["status"]}
            )
        dt = datetime.fromisoformat(a["appointment_datetime"])
        if str(e["AppointmentDate"])[:10] != dt.date().isoformat() or \
           str(e["AppointmentTime"])[:5] != dt.strftime("%H:%M"):
            findings["DATETIME_DRIFT"].append(
                {"ref": ref,
                 "excel": f"{e['AppointmentDate']} {e['AppointmentTime']}",
                 "hms": dt.isoformat(timespec="minutes")}
            )

    for ref in excel:
        if ref not in appts:
            # Never auto-deleted. This may be the only record of a lost booking.
            findings["ORPHAN_IN_EXCEL"].append(ref)

    return findings


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Manage the Excel mirror")
    ap.add_argument("command", choices=["init", "sync", "reconcile"])
    ap.add_argument("--base", default="http://localhost:5000", help="HMS base URL")
    ap.add_argument("--key", default="dev-api-key-change-me", help="HMS API key")
    ap.add_argument("--path", default=str(EXCEL_PATH))
    args = ap.parse_args(argv)
    path = Path(args.path)

    if args.command == "init":
        print(f"created {init(path)}")
        return 0

    if args.command == "sync":
        r = sync(args.base, args.key, path)
        print(f"HMS rows {r['hms_rows']} | already in Excel {r['already_present']} | added {r['added']}")
        return 0

    r = reconcile(args.base, args.key, path)
    total = sum(len(v) for v in r.values())
    print(f"Reconciliation: {total} finding(s)")
    for kind, items in r.items():
        if not items:
            continue
        auto = "auto-fixable" if kind != "ORPHAN_IN_EXCEL" else "ESCALATE — never auto-deleted"
        print(f"  {kind:18} {len(items):4}  ({auto})")
        for item in items[:5]:
            print(f"      {item}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
